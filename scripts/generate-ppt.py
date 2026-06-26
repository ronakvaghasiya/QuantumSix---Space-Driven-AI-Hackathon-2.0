#!/usr/bin/env python3
"""
Generate RepoPilot AI PPT aligned with demo video (demo-narration.json).
UI styling matches the app: dark title slides, green accents, framed screenshots.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
SHOTS = ROOT / "docs" / "screenshots"
CACHE = ROOT / "docs" / "ppt-cache"
NARRATION = SCRIPTS / "demo-narration.json"
OUT = ROOT / "docs" / "RepoPilot_AI_Presentation.pptx"
OUT_FULL = ROOT / "docs" / "RepoPilot_AI_Full_Presentation.pptx"

# App theme (matches video + AuthShell dark green)
GREEN = RGBColor(34, 197, 94)
GREEN_DARK = RGBColor(22, 163, 74)
DARK = RGBColor(15, 23, 42)
DARK_CARD = RGBColor(30, 41, 59)
TEXT = RGBColor(30, 41, 59)
MUTED = RGBColor(100, 116, 139)
WHITE = RGBColor(255, 255, 255)
LIGHT_BG = RGBColor(248, 250, 252)
BORDER = RGBColor(226, 232, 240)

SLIDE_W = Inches(13.333)  # 16:9 widescreen
SLIDE_H = Inches(7.5)

TAGLINE = "From a Prompt to Production"
ORG = "QuantumSix · Space-Driven AI Hackathon 2.0"


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def footer(slide, text: str, slide_no: int | None = None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), SLIDE_H - Inches(0.38), SLIDE_W, Inches(0.38))
    bar.fill.solid()
    bar.fill.fore_color.rgb = DARK
    bar.line.fill.background()

    left = slide.shapes.add_textbox(Inches(0.5), SLIDE_H - Inches(0.34), Inches(8), Inches(0.28))
    p = left.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(9)
    p.font.color.rgb = MUTED

    if slide_no is not None:
        right = slide.shapes.add_textbox(SLIDE_W - Inches(1.2), SLIDE_H - Inches(0.34), Inches(0.8), Inches(0.28))
        p2 = right.text_frame.paragraphs[0]
        p2.text = str(slide_no)
        p2.font.size = Pt(9)
        p2.font.color.rgb = GREEN
        p2.alignment = PP_ALIGN.RIGHT


def top_accent(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(0.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()


def badge(slide, text: str, left=Inches(0.55), top=Inches(0.35)):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.8), Inches(0.32))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(220, 252, 231)
    box.line.color.rgb = GREEN
    box.line.width = Pt(0.75)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text.upper()
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = GREEN_DARK
    p.alignment = PP_ALIGN.CENTER


def slide_title(slide, title: str, top=Inches(0.78)):
    box = slide.shapes.add_textbox(Inches(0.55), top, Inches(12), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = TEXT


def bullets(slide, items: list[str], left=Inches(0.55), top=Inches(1.55), width=Inches(12), size=13):
    box = slide.shapes.add_textbox(left, top, width, Inches(5.2))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = TEXT
        p.space_after = Pt(8)
        p.level = 0


def narration_to_bullets(text: str, max_items: int = 4) -> list[str]:
    """Split video narration into short presenter bullets."""
    cleaned = re.sub(r"\s+", " ", text.strip())
    parts = re.split(r"(?<=[.!?])\s+", cleaned)
    out = []
    for part in parts:
        part = part.strip()
        if len(part) < 12:
            continue
        if len(part) > 110:
            part = part[:107] + "..."
        out.append(f"• {part}")
        if len(out) >= max_items:
            break
    return out or [f"• {cleaned[:100]}..."]


def prepare_screenshot(src: Path, max_h: int = 1200) -> Path:
    """Crop very tall screenshots to top region for readable PPT layout."""
    CACHE.mkdir(parents=True, exist_ok=True)
    out = CACHE / f"{src.stem}-ppt.png"
    try:
        from PIL import Image
        img = Image.open(src)
        w, h = img.size
        if h > max_h:
            img = img.crop((0, 0, w, max_h))
        if w > 1920:
            ratio = 1920 / w
            img = img.resize((1920, int(h * ratio)), Image.Resampling.LANCZOS)
        img.save(out, optimize=True)
        return out
    except Exception:
        return src


def add_browser_frame(slide, img_path: Path, left: float, top: float, width: float) -> None:
    """Place screenshot inside a browser-style frame."""
    frame_l = Inches(left)
    frame_t = Inches(top)
    frame_w = Inches(width)
    chrome_h = Inches(0.38)

    # Shadow card
    shadow = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, frame_l + Inches(0.06), frame_t + Inches(0.06), frame_w, Inches(5.55))
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = RGBColor(203, 213, 225)
    shadow.line.fill.background()

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, frame_l, frame_t, frame_w, Inches(5.55))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = BORDER
    card.line.width = Pt(1)

    # Chrome bar
    chrome = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, frame_l, frame_t, frame_w, chrome_h)
    chrome.fill.solid()
    chrome.fill.fore_color.rgb = DARK_CARD
    chrome.line.fill.background()

    for i, color in enumerate([RGBColor(239, 68, 68), RGBColor(234, 179, 8), RGBColor(34, 197, 94)]):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, frame_l + Inches(0.15 + i * 0.18), frame_t + Inches(0.12), Inches(0.14), Inches(0.14))
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()

    url_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, frame_l + Inches(0.75), frame_t + Inches(0.08), frame_w - Inches(1.1), Inches(0.22))
    url_box.fill.solid()
    url_box.fill.fore_color.rgb = RGBColor(51, 65, 85)
    url_box.line.fill.background()
    url_tf = url_box.text_frame
    url_tf.clear()
    up = url_tf.paragraphs[0]
    up.text = "localhost:3100"
    up.font.size = Pt(8)
    up.font.color.rgb = MUTED
    up.alignment = PP_ALIGN.CENTER

    pic_top = frame_t + chrome_h + Inches(0.05)
    pic_h_avail = Inches(5.55) - chrome_h - Inches(0.1)
    pic = slide.shapes.add_picture(str(img_path), frame_l + Inches(0.08), pic_top, width=frame_w - Inches(0.16))
    if pic.height > pic_h_avail:
        ratio = pic.width / pic.height
        pic.height = int(pic_h_avail)
        pic.width = int(pic.height * ratio)


def add_title_slide(prs, title: str, subtitle: str, badge_text: str = ORG):
    slide = blank_slide(prs)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = DARK
    top_accent(slide)

    glow = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.5), Inches(-1.5), Inches(5), Inches(5))
    glow.fill.solid()
    glow.fill.fore_color.rgb = RGBColor(22, 101, 52)
    glow.fill.transparency = 0.75
    glow.line.fill.background()

    b = slide.shapes.add_textbox(Inches(0.9), Inches(0.9), Inches(10), Inches(0.35))
    bp = b.text_frame.paragraphs[0]
    bp.text = badge_text.upper()
    bp.font.size = Pt(10)
    bp.font.bold = True
    bp.font.color.rgb = GREEN

    t = slide.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(11), Inches(1.1))
    tp = t.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(44)
    tp.font.bold = True
    tp.font.color.rgb = WHITE

    s = slide.shapes.add_textbox(Inches(0.9), Inches(2.85), Inches(10), Inches(2.5))
    stf = s.text_frame
    stf.word_wrap = True
    for i, line in enumerate(subtitle.split("\n")):
        sp = stf.paragraphs[0] if i == 0 else stf.add_paragraph()
        sp.text = line
        sp.font.size = Pt(18 if i == 0 else 14)
        sp.font.color.rgb = MUTED if i > 0 else WHITE
        sp.space_after = Pt(10)

    footer(slide, TAGLINE)
    return slide


def add_section_slide(prs, number: str, title: str, subtitle: str):
    slide = blank_slide(prs)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = DARK
    top_accent(slide)

    num = slide.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(3), Inches(1))
    np = num.text_frame.paragraphs[0]
    np.text = number
    np.font.size = Pt(64)
    np.font.bold = True
    np.font.color.rgb = GREEN

    t = slide.shapes.add_textbox(Inches(0.9), Inches(3.3), Inches(11), Inches(0.9))
    tp = t.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(36)
    tp.font.bold = True
    tp.font.color.rgb = WHITE

    s = slide.shapes.add_textbox(Inches(0.9), Inches(4.3), Inches(10), Inches(1))
    sp = s.text_frame.paragraphs[0]
    sp.text = subtitle
    sp.font.size = Pt(16)
    sp.font.color.rgb = MUTED

    footer(slide, "RepoPilot AI")
    return slide


def add_content_slide(prs, section: str, title: str, items: list[str], note: str = ""):
    slide = blank_slide(prs)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT_BG
    top_accent(slide)
    badge(slide, section)
    slide_title(slide, title)
    bullets(slide, items, size=14)
    if note:
        n = slide.shapes.add_textbox(Inches(0.55), Inches(6.55), Inches(12), Inches(0.4))
        np = n.text_frame.paragraphs[0]
        np.text = note
        np.font.size = Pt(10)
        np.font.italic = True
        np.font.color.rgb = GREEN_DARK
    footer(slide, TAGLINE)
    return slide


def add_table_slide(prs, section: str, title: str, headers: list[str], rows: list[list[str]], col_widths=None):
    slide = blank_slide(prs)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT_BG
    top_accent(slide)
    badge(slide, section)
    slide_title(slide, title)

    cols = len(headers)
    rows_n = len(rows) + 1
    tbl_h = min(Inches(0.38) * rows_n, Inches(5.2))
    shape = slide.shapes.add_table(rows_n, cols, Inches(0.5), Inches(1.55), Inches(12.3), tbl_h)
    table = shape.table

    if col_widths:
        for j, w in enumerate(col_widths):
            table.columns[j].width = Inches(w)

    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = GREEN
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(10)
            p.font.color.rgb = WHITE

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = str(val)
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9)
                p.font.color.rgb = TEXT

    footer(slide, TAGLINE)
    return slide


def add_ui_slide(prs, scene_no: int, section: str, title: str, narration: str, image_name: str | None):
    """Video-aligned screenshot slide — left bullets, right browser frame."""
    slide = blank_slide(prs)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT_BG
    top_accent(slide)

    badge(slide, f"{section} · Screen {scene_no:02d}")
    slide_title(slide, title)

    # Green accent line under title
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.42), Inches(1.8), Inches(0.05))
    line.fill.solid()
    line.fill.fore_color.rgb = GREEN
    line.line.fill.background()

    bullet_items = narration_to_bullets(narration, 5)
    bullets(slide, bullet_items, left=Inches(0.55), top=Inches(1.65), width=Inches(4.5), size=11)

    if image_name:
        src = SHOTS / image_name
        if src.exists():
            img = prepare_screenshot(src)
            add_browser_frame(slide, img, left=5.35, top=1.35, width=7.45)
        else:
            bullets(slide, [f"Screenshot missing: {image_name}"], left=Inches(5.5), top=Inches(2), width=Inches(6), size=12)

    footer(slide, TAGLINE, scene_no)
    return slide


def load_narration() -> dict:
    return json.loads(NARRATION.read_text())


def scene_by_id(scenes: list, sid: str) -> dict:
    for s in scenes:
        if s["id"] == sid:
            return s
    raise KeyError(sid)


def main():
    cfg = load_narration()
    scenes = cfg["scenes"]

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    intro = scene_by_id(scenes, "intro")
    outro = scene_by_id(scenes, "outro")

    # ── OPENING (matches video intro) ──
    add_title_slide(
        prs,
        "RepoPilot AI",
        f"{TAGLINE}\n\nSDLC Engineering Copilot for GitLab Teams\n"
        "Governed ticket-to-merge-request pipeline with human approval gates",
    )

    add_content_slide(prs, "Agenda", "Presentation Flow (Same as Demo Video)", [
        "1. What is RepoPilot AI — problem, solution, 9-step pipeline",
        "2. Live UI walkthrough — Login → Dashboard → Projects → Tasks",
        "3. Task BB-007 end-to-end — Analysis, Code Diff, Validation, GitLab MR",
        "4. Human approval gates — DEMO-004 & DEMO-005",
        "5. Reports, Settings, architecture & competitive positioning",
    ])

    add_content_slide(
        prs, "Overview", "What is RepoPilot AI?",
        narration_to_bullets(intro["narration"], 6),
        note=ORG,
    )

    add_table_slide(prs, "Problem", "Software Delivery is Fragmented", ["Pain Point", "Impact"], [
        ["Manual ticket → code → MR cycle", "Slow delivery, context switching"],
        ["Legacy repos hard to navigate", "Wrong files changed, regressions"],
        ["No governance on AI code changes", "Trust issues, audit gaps"],
        ["Scattered tools (Jira, IDE, GitLab, CI)", "No single pane of glass"],
    ], [5.5, 6.8])

    add_content_slide(prs, "Solution", "One Dashboard — Full SDLC", [
        "• Tickets enter as CSV — task_id + requirement columns teams already use",
        "• AI runs analysis, repo search, test generation, code gen, and validation",
        "• Humans approve at TWO gates — before code is written and before MR is created",
        "• Output: real GitLab Merge Request with lint, build, and security checks passed",
        "• Engineering managers get KPIs, pipeline status, and audit trail in one place",
    ])

    pipeline = scene_by_id(scenes, "09-pipeline")
    add_table_slide(prs, "Pipeline", "SDLC Pipeline — 9 Steps", ["#", "Step", "What Happens"], [
        ["1", "Requirement Analysis", "AI generates acceptance criteria & user stories"],
        ["2", "Repository Intelligence", "Qdrant semantic search finds impacted files"],
        ["3", "Impact Analysis", "Regression areas & API dependencies identified"],
        ["4", "Test Generation", "QA test cases & Playwright specs created"],
        ["5", "Human Approval", "Analysis gate — approve before any code is written"],
        ["6", "Code Generation", "Deterministic replace or LLM file edit"],
        ["7", "Human Approval", "Code diff review gate before MR"],
        ["8", "Validation + QA", "ESLint, build, Playwright, security scan"],
        ["9", "GitLab MR", "Branch, commit, push, merge request created"],
    ], [0.6, 3.8, 7.9])

    add_content_slide(prs, "Pipeline", "Task Status Flow & Demo Tickets", [
        "pending → analyzing → analysis_approval_required → generating_code",
        "→ code_approval_required → validating → security_scan → creating_pr → pr_created",
        "",
        "• BB-005 (demo upload): change label \"Sides Preview\" → \"Sides Preview 123\"",
        "• BB-007 (live example): completed pipeline with GitLab MR !868",
        "• DEMO-004: waiting at analysis approval gate",
        "• DEMO-005: waiting at code approval gate",
    ])

    add_section_slide(prs, "01", "Live UI Walkthrough", "Same screens & flow as the demo video")

    # ── UI SLIDES (video order) ──
    ui_scene_ids = [
        "01-login", "02-dashboard", "03-projects", "04-project-detail",
        "05-knowledge-graph", "06-tasks", "07-upload", "08-overview",
        "09-pipeline", "10-analysis", "11-tests", "12-diff",
        "13-validation", "14-pr", "15-audit", "16-analysis-gate",
        "17-code-gate", "18-reports", "19-settings",
    ]

    for i, sid in enumerate(ui_scene_ids, 1):
        sc = scene_by_id(scenes, sid)
        add_ui_slide(
            prs, i, "UI Demo", sc["title"],
            sc["narration"], sc.get("image"),
        )

    add_section_slide(prs, "02", "Platform Deep Dive", "Architecture, stack & positioning")

    add_content_slide(prs, "Architecture", "System Architecture", [
        "Browser (Next.js :3100) → REST API + JWT → NestJS Backend (:3001)",
        "  ├── TaskPipelineService — orchestrates all 9 pipeline steps",
        "  ├── Projects / Indexing / GitLab / Reports modules",
        "  └── PostgreSQL — tasks, timeline, diffs, merge requests",
        "       ├── Qdrant — vector embeddings for semantic file search",
        "       ├── GitLab API — clone, branch, commit, MR create/merge",
        "       └── Local git clone — real commits pushed to GitLab remote",
    ])

    add_table_slide(prs, "Tech Stack", "Technology Stack", ["Layer", "Technology", "Purpose"], [
        ["Frontend", "Next.js 14, TypeScript, MUI", "Dashboard UI"],
        ["Backend", "NestJS, TypeORM", "API + pipeline orchestration"],
        ["Database", "PostgreSQL 16", "Persistent data"],
        ["Vector DB", "Qdrant", "Semantic file search"],
        ["AI (default)", "Qwen 2.5 + MiniLM embeddings", "Chat + search"],
        ["VCS", "GitLab REST API", "Clone, MR, merge"],
        ["Validation", "ESLint, Prettier, Playwright", "Quality gate"],
    ], [2, 4.5, 5.8])

    add_table_slide(prs, "Competitive", "RepoPilot vs Market", ["Tool", "RepoPilot Advantage"], [
        ["GitHub Copilot", "Full SDLC loop — not just IDE autocomplete"],
        ["Cursor", "Team governance + MR — not single-developer IDE"],
        ["Devin", "Human approval gates + full audit trail"],
        ["n8n / Zapier", "Deep repo intelligence + verified code changes"],
    ], [3.2, 9.1])

    add_table_slide(prs, "Demo Guide", "Golden Demo Path (2 Minutes)", ["Time", "Action"], [
        ["0:00", "Dashboard — KPIs & pipeline status"],
        ["0:20", "Projects — DesigntoolReact indexed repo"],
        ["0:40", "Tasks — upload BB-005.csv"],
        ["1:00", "Approve analysis — show impacted files"],
        ["1:20", "Code Diff — Sides Preview → Sides Preview 123"],
        ["1:40", "Approve code → show GitLab MR !868"],
        ["2:00", "Close: From a Prompt to Production"],
    ], [1.8, 10.5])

    add_title_slide(
        prs,
        "Thank You",
        f"{TAGLINE}\n\n"
        "Cursor helps one developer type faster.\n"
        "RepoPilot governs the entire change for GitLab teams.\n\n"
        "Dashboard: localhost:3100  |  API: localhost:3001/docs\n\n"
        f"{ORG}",
        "Questions?",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    prs.save(str(OUT_FULL))

    import shutil
    dl = Path.home() / "Downloads" / "RepoPilot_AI_Presentation.pptx"
    try:
        shutil.copy2(OUT, dl)
        print(f"Copied: {dl}")
    except OSError:
        pass

    print(f"Created: {OUT}")
    print(f"Created: {OUT_FULL}")
    print(f"Total slides: {len(prs.slides)}")
    print("Aligned with: scripts/demo-narration.json (demo video)")


if __name__ == "__main__":
    main()
